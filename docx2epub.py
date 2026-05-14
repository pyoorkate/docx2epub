import sys
import os
import re
from docx import Document
from ebooklib import epub
import uuid

def sanitize_style_name(name):
    """
    Converts 'List Paragraph' to 'List_Paragraph'
    and removes non-alphanumeric characters for valid CSS class names.
    """
    # Replace spaces with underscores
    clean = name.replace(' ', '_')
    # Remove anything that isn't alphanumeric or underscore
    clean = re.sub(r'[^a-zA-Z0-9_]', '', clean)
    return clean

def create_css():
    """Defines the stylesheet."""
    style = """
    @namespace epub "http://www.idpf.org/2007/ops";

    body {
        font-family: serif;
        margin: 5%;
    }

    p {
        font-family: serif;
        text-align: justify;
        line-height: 1.2;
        margin-bottom: 1em;
    }

    h1 {
        font-family: sans-serif;
        text-align: center;
        margin-bottom: 1.5em;
    }

    /* Standard formatting classes */
    .bold { font-weight: bold; }
    .italic { font-style: italic; }
    .underline { text-decoration: underline; }
    .strikethrough { text-decoration: line-through; }

    /* Default styling for generated divs (fallback) */
    div {
        margin-bottom: 1em;
    }
    """
    return epub.EpubItem(uid="style_nav", file_name="style/nav.css", media_type="text/css", content=style)

def get_run_html(run):
    """Converts a docx run to HTML text with formatting."""
    text = run.text
    if not text:
        return ""

    # Escape HTML special characters
    text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    if run.bold:
        text = f'<div class="bold">{text}</div>'
    if run.italic:
        text = f'<div class="italic">{text}</div>'
    if run.underline:
        text = f'<div class="underline">{text}</div>'
    if run.font.strike:
        text = f'<div class="strikethrough">{text}</div>'

    return text

def create_chapter_item(title, content, file_name):
    """Creates an ebooklib chapter item."""
    chapter = epub.EpubHtml(title=title, file_name=file_name, lang='en')

    # Trust that the content buffer already contains the H1 if needed.
    # Reliant on either reliably using H1 or using docx cleaning script.
    chapter.content = content

    chapter.add_link(href="style/nav.css", rel="stylesheet", type="text/css")
    return chapter

def convert_docx_to_epub(docx_path, output_path):
    # 1. Load Document
    if not os.path.exists(docx_path):
        print(f"Error: File {docx_path} not found.")
        return

    doc = Document(docx_path)
    book = epub.EpubBook()

    # 2. Metadata
    book.set_identifier(str(uuid.uuid4()))
    book.set_title(os.path.basename(docx_path).replace('.docx', ''))
    book.set_language('en')
    book.add_author('Unknown Author')

    # 3. Add CSS
    css_item = create_css()
    book.add_item(css_item)

    # 4. Parsing Variables
    chapters = []
    current_chapter_content = []
    current_chapter_title = "Preface"

    print("Parsing document...")

    # 5. Parse Paragraphs
    for para in doc.paragraphs:
        # Skip completely empty paragraphs
        if not para.text.strip():
            continue

        style_name = para.style.name

        # Determine HTML tag and Class
        tag = 'div'
        css_class = ""

        if style_name == 'Normal':
            tag = 'p'
        elif style_name == 'Heading 1':
            tag = 'h1'
        else:
            # For all other styles, use a div with the sanitized style name
            tag = 'div'
            clean_style = sanitize_style_name(style_name)
            css_class = f' class="{clean_style}"'

        # Generate HTML for the content
        inner_html = "".join([get_run_html(run) for run in para.runs])

        # --- LOGIC FOR SPLITTING CHAPTERS ---
        if tag == 'h1':
            # If we have content in the buffer, save it as the PREVIOUS chapter
            if current_chapter_content:
                chap_file = f"chapter_{len(chapters)}.xhtml"
                full_html = "".join(current_chapter_content)
                chap_item = create_chapter_item(current_chapter_title, full_html, chap_file)
                chapters.append(chap_item)

            # Start New Chapter
            current_chapter_title = para.text.strip()
            current_chapter_content = []

            # Add the Heading 1 to the NEW chapter's content
            current_chapter_content.append(f"<h1>{inner_html}</h1>")

        else:
            # Just append to current buffer
            current_chapter_content.append(f"<{tag}{css_class}>{inner_html}</{tag}>")

    # 6. Flush the final chapter
    if current_chapter_content:
        chap_file = f"chapter_{len(chapters)}.xhtml"
        full_html = "".join(current_chapter_content)
        chap_item = create_chapter_item(current_chapter_title, full_html, chap_file)
        chapters.append(chap_item)

    # 7. Create Placeholder Pages

    # Helper to make simple pages
    def make_page(title, filename, body_text):
        page = epub.EpubHtml(title=title, file_name=filename)
        # Use a simple div for centering placeholders
        page.content = f'<div style="text-align:center; margin-top:30%;"><h1>{title}</h1><p>{body_text}</p></div>'
        page.add_link(href="style/nav.css", rel="stylesheet", type="text/css")
        return page

    cover_page = make_page('Cover Page', 'cover.xhtml', 'cover page')
    title_page = make_page('Title Page', 'title.xhtml', 'title page')
    author_page = make_page('About the Author', 'about_author.xhtml', 'about the author')
    press_page = make_page('About the Press', 'about_press.xhtml', 'about the press')
    copyright_page = make_page('Copyright Information', 'copyright.xhtml', 'insert copyright information')

    # 8. Add Items to Book (Order here doesn't affect display, only internal storage)
    book.add_item(cover_page)
    book.add_item(title_page)
    for chap in chapters:
        book.add_item(chap)
    book.add_item(author_page)
    book.add_item(press_page)
    book.add_item(copyright_page)

    # 9. Define Table of Contents (TOC)
    book.toc = chapters

    # 10. Define Spine (The actual reading order)
    spine_list = ['nav', cover_page, title_page] + chapters + [author_page, press_page, copyright_page]
    book.spine = spine_list

    # 11. Navigation Files
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())

    # 12. Write Output
    print(f"Writing {output_path}...")
    epub.write_epub(output_path, book, {})
    print("Done!")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python docx2epub_v2.py input_file.docx [output_file.epub]")
    else:
        input_file = sys.argv[1]
        output_file = sys.argv[2] if len(sys.argv) > 2 else input_file.replace('.docx', '.epub')
        convert_docx_to_epub(input_file, output_file)
